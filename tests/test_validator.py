"""Slice 5: validation reports the right problems, and the screen lets you fix them."""
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from src import ai
from src.mapper import name_match
from src.master_reader import read_master
from src.questions import Answer, build_queue, resolve_all
from src.result import build_rows
from src.template_reader import read_template
from src.validator import CONDITIONAL_MISSING, INVALID, REQUIRED_MISSING, blocking, validate

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "sample"
TEMPLATE, MASTER = SAMPLE / "RING (6).xlsm", SAMPLE / "Master_Arias_Ring_LOT3_v2.xlsx"
KEY = SAMPLE / "Amazon_Ring upd_ (1).xlsm"
pytestmark = pytest.mark.skipif(not (TEMPLATE.exists() and MASTER.exists() and KEY.exists()),
                                reason="sample/ files not present")


@pytest.fixture(scope="module")
def ctx():
    info, master = read_template(TEMPLATE.read_bytes()), read_master(MASTER.read_bytes())
    mapping = name_match(info.columns, master.columns)
    mapping[next(c.index for c in info.columns if c.label == "SKU")] = "Sku Code"
    return info, master, mapping


def test_skipped_required_and_bad_dropdown_are_reported(ctx):
    info, master, mapping = ctx
    rows = build_rows(info, master, mapping)
    issues = validate(info, rows)
    required = {i.col.label: i for i in issues if i.severity == REQUIRED_MISSING}
    assert "Item Name" in required or "Product Description" in required     # unmapped required columns
    assert len(required["Product Description"].rows) == len(master)         # row counts are reported
    assert all(r >= info.data_row for r in required["Product Description"].rows)
    assert blocking(issues) == [i for i in issues if i.severity == REQUIRED_MISSING]
    # a value outside a dropdown list is flagged (as a "check this" item), with the offending value
    drop = next(c for c in info.columns if c.has_dropdown and c.required)
    bad_rows = build_rows(info, master, mapping, {drop.index: "not a real value"})
    flagged = [i for i in validate(info, bad_rows) if i.severity == INVALID and i.col.index == drop.index]
    assert flagged and flagged[0].examples == ["not a real value"]


def test_complete_required_columns_report_no_blocking_issue(ctx):
    info, master, mapping = ctx
    answers = {c.index: Answer("const", value="x" if not c.has_dropdown else c.allowed[0])
               for c in info.required_columns() if c.index not in mapping}
    parentage = next(c.index for c in info.columns if c.label == "Parentage Level")
    answers[parentage] = Answer("lookup", by="CATEGORY", table={"PARENT": "Parent", "RING": "Child"})
    rows = build_rows(info, master, mapping, resolve_all(answers, master))
    assert blocking(validate(info, rows)) == []


def test_hand_filled_file_is_nearly_clean():
    """The answer key is a file Amazon accepts; the validator should find (almost) nothing blocking in it."""
    info = read_template(TEMPLATE.read_bytes())
    ws = openpyxl.load_workbook(KEY, data_only=True)["Template"]
    rows = []
    for r in range(info.data_row, ws.max_row + 1):
        row = {c.index: ws.cell(r, c.index).value for c in info.columns if ws.cell(r, c.index).value not in (None, "")}
        if row:
            rows.append(row)
    assert len(rows) == 608
    must = blocking(validate(info, rows))
    assert must == []        # parent rows are exempt from the product id check


class FakeClient:
    def __init__(self, reply):
        self.reply, self.calls = reply, 0
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kw):
        self.calls += 1
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=self.reply)])


def test_answer_check_parses_and_fails_plainly(ctx):
    info, *_ = ctx
    col = next(c for c in info.columns if c.label == "Item Name")
    good = FakeClient('{"results": [{"value": "Ring", "ok": false, "reason": "too short"}]}')
    assert ai.check_answers(good, col, ["Ring"]) == [("Ring", False, "too short")]
    with pytest.raises(ai.AIUnavailable, match="could not be completed"):
        ai.check_answers(FakeClient("nope"), col, ["Ring"])
    with pytest.raises(ai.AIUnavailable, match="API key"):
        ai.check_answers(None, col, ["Ring"])


def test_validation_screen_flags_then_fix_reopens_the_question(ctx):
    info, master, mapping = ctx
    tb = TEMPLATE.read_bytes()
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=90)
    at.session_state["master"], at.session_state["template"] = master, info
    at.session_state["template_bytes"], at.session_state["template_name"] = tb, "RING (6).xlsm"
    at.run()
    assert not at.exception
    assert any("Required columns are still empty" in e.value or "required columns are still empty" in e.value
               for e in at.error)
    fix = next(b for b in at.button if b.label == "Fix this")
    fix.click()
    at.run()
    assert not at.exception
    queue = build_queue(info, at.session_state["mapping"], master, False, at.session_state["q_extra"])
    assert queue[at.session_state["q_index"]].required        # jumped to a required column's question
