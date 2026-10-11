"""Slice 7: the built-in listing rules, checked against the hand-filled Amazon file."""
import io
from pathlib import Path

import openpyxl
import pytest

from src.mapper import name_match
from src.master_reader import read_master
from src.questions import build_queue
from src.result import build_rows
from src.rules import apply_rules, find_column, load_rules
from src.template_reader import read_template
from src.validator import blocking, validate
from src.xlsm_writer import fill_template

SAMPLE = Path(__file__).resolve().parent.parent / "sample"
TEMPLATE, MASTER = SAMPLE / "RING (6).xlsm", SAMPLE / "Master_Arias_Ring_LOT3_v2.xlsx"
KEY = SAMPLE / "Amazon_Ring upd_ (1).xlsm"
pytestmark = pytest.mark.skipif(not (TEMPLATE.exists() and MASTER.exists() and KEY.exists()),
                                reason="sample/ files not present")


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


@pytest.fixture(scope="module")
def ctx():
    info, master = read_template(TEMPLATE.read_bytes()), read_master(MASTER.read_bytes())
    res = apply_rules(info, master, load_rules())
    key = openpyxl.load_workbook(KEY, data_only=True)["Template"]
    return info, master, res, key


def test_parent_child_structure(ctx):
    info, master, res, _ = ctx
    parentage = find_column(info, "Parentage Level").index
    assert [v for v in res.values[parentage]].count("parent") == 32 == sum(res.is_parent)
    assert res.values[parentage].count("child") == 576
    sku = find_column(info, "SKU").index
    parent_sku = find_column(info, "Parent SKU").index
    child_rows = [i for i, p in enumerate(res.is_parent) if not p]
    parents = {master.iloc[i]["CODE"]: master.iloc[i]["Sku Code"] for i, p in enumerate(res.is_parent) if p}
    assert all(res.values[parent_sku][i] == parents[master.iloc[i]["CODE"]] for i in child_rows)
    assert all(res.values[parent_sku][i] is None for i, p in enumerate(res.is_parent) if p)   # parents have none


def test_translations_and_derivations(ctx):
    info, master, res, _ = ctx
    get = lambda label, i: res.values[find_column(info, label).index][i]
    child = next(i for i, p in enumerate(res.is_parent) if not p)
    assert master.iloc[child]["Certificate Type"] == "IGI Certificate" and get("Certificate Type", child) == "IGI"
    assert get("Stone Shape", child) == "Round Brilliant"
    assert get("Target Gender", child) == "female" and get("Department Name", child) == "Womens"
    assert get("Product Type", child) == "RING"
    assert get("Metal Stamp#1", child) == "14K"
    assert get("Material#1", child) == f"14k {master.iloc[child]['Color']}"
    assert get("Color Map", child) == master.iloc[child]["Color"].replace(" Gold", "")
    assert "Rings (2152567031)" in get("Recommended Browse Nodes", child) and "Women" in get("Recommended Browse Nodes", child)


def test_parents_inherit_from_first_child_but_child_only_stay_blank(ctx):
    info, master, res, _ = ctx
    parent = next(i for i, p in enumerate(res.is_parent) if p)
    assert master.iloc[parent]["Certificate Type"] is None or str(master.iloc[parent]["Certificate Type"]) == "nan" \
        or master.iloc[parent]["Certificate Type"] != master.iloc[parent]["Certificate Type"]
    assert res.values[find_column(info, "Certificate Type").index][parent] == "IGI"       # inherited
    for label in ("Ring Size", "Metal Weight", "Main Image URL", "Material#1", "Parent SKU"):
        idx = find_column(info, label).index
        assert idx in res.child_only and res.values[idx][parent] is None
    # Product Id Type is deliberately not inherited: parents carry no product id
    pid = find_column(info, "Product Id Type").index
    assert res.values[pid][0] is None


def test_rules_reproduce_the_hand_filled_file(ctx):
    info, master, res, key = ctx
    differing, total, matched = {}, 0, 0
    for idx, vals in res.values.items():
        for i, v in enumerate(vals):
            k = norm(key.cell(info.data_row + i, idx).value)
            total += 1
            if k == norm(v):
                matched += 1
            else:
                differing[info.columns[idx - 1].label] = differing.get(info.columns[idx - 1].label, 0) + 1
    print(f"\nrules: {len(res.values)} columns, {matched}/{total} cells match ({matched / total:.1%}); "
          f"differing: {differing}")
    assert len(res.values) >= 45 and matched / total >= 0.92
    # the differences are the editorial picks in the hand-filled file (images), one unit cell, and the variation
    # theme: the rules use COLOR/RING_SIZE because the children differ by colour AND ring size (the AI review
    # flagged the hand-filled file's plain COLOR, and the template says all theme attributes must be populated)
    assert set(differing) <= {"Main Image URL", "Other Image URL", "Item Weight Unit", "Variation Theme Name"}


def test_rule_values_fit_the_templates_dropdowns(ctx):
    assert ctx[2].problems == []
    assert ctx[2].skipped == []


def test_questions_shrink_when_rules_are_on(ctx):
    info, master, res, _ = ctx
    mapping = name_match(info.columns, master.columns)
    without = build_queue(info, mapping, master)
    with_rules = build_queue(info, mapping, master, covered=res.covered())
    assert len(with_rules) < len(without)
    assert not {c.index for c in with_rules} & res.covered()
    print(f"\nquestions: {len(without)} without rules -> {len(with_rules)} with rules")


def test_rules_beat_raw_master_values_and_answers_only_fill_gaps(ctx):
    info, master, res, _ = ctx
    mapping = name_match(info.columns, master.columns)
    cert = find_column(info, "Certificate Type")
    mapping[cert.index] = "Certificate Type"                    # the raw master value is "IGI Certificate"
    rows = build_rows(info, master, mapping, {cert.index: "WRONG"}, res.values)
    child = next(i for i, p in enumerate(res.is_parent) if not p)
    assert rows[child][cert.index] == "IGI"                      # the rule replaces the raw value; the answer only fills gaps


def test_end_to_end_with_rules_leaves_few_blocking_problems(ctx):
    info, master, res, _ = ctx
    mapping = name_match(info.columns, master.columns)
    rows = build_rows(info, master, mapping, None, res.values)
    out = fill_template(TEMPLATE.read_bytes(), info.data_row, rows)
    ws = openpyxl.load_workbook(io.BytesIO(out), data_only=True)["Template"]
    sku = find_column(info, "SKU").index
    assert ws.cell(info.data_row, sku).value == master.iloc[0]["Sku Code"]                # SKU comes from the rule
    assert ws.cell(info.data_row, find_column(info, "Parentage Level").index).value == "parent"
    must = blocking(validate(info, rows, res.child_only))
    print("\nstill required and empty without any answers:", [(i.col.label, len(i.rows)) for i in must])
    assert len(must) <= 2
