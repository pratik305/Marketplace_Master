"""Slice 2: the writer fills the Template sheet and leaves everything else untouched."""
import io
import zipfile
from pathlib import Path

import openpyxl
import pytest

from src.mapper import name_match
from src.master_reader import read_master
from src.result import build_rows
from src.template_reader import read_template
from src.xlsm_writer import fill_template, sheet_path

SAMPLE = Path(__file__).resolve().parent.parent / "sample"
TEMPLATE = SAMPLE / "RING (6).xlsm"
MASTER = SAMPLE / "Master_Arias_Ring_LOT3_v2.xlsx"
pytestmark = pytest.mark.skipif(not (TEMPLATE.exists() and MASTER.exists()), reason="sample/ files not present")


@pytest.fixture(scope="module")
def produced():
    tb = TEMPLATE.read_bytes()
    info = read_template(tb)
    master = read_master(MASTER.read_bytes())
    mapping = name_match(info.columns, master.columns)
    # Plain name matching can't pair "SKU" with "Sku Code"; the AI mapping (slice 3) does that.
    mapping[next(c.index for c in info.columns if c.label == "SKU")] = "Sku Code"
    rows = build_rows(info, master, mapping)
    return tb, info, master, mapping, rows, fill_template(tb, info.data_row, rows)


def test_only_template_sheet_changes(produced):
    tb, *_, out = produced
    zin, zout = zipfile.ZipFile(io.BytesIO(tb)), zipfile.ZipFile(io.BytesIO(out))
    changed = sheet_path(zin, "Template")
    assert [i.filename for i in zin.infolist()] == [i.filename for i in zout.infolist()]
    for name in zin.namelist():
        if name != changed:
            assert zin.read(name) == zout.read(name), f"{name} was modified"
    assert zin.read(changed) != zout.read(changed)


def test_dropdowns_and_images_survive(produced):
    tb, *_, out = produced
    sheet = sheet_path(zipfile.ZipFile(io.BytesIO(tb)), "Template")
    a, b = zipfile.ZipFile(io.BytesIO(tb)).read(sheet).decode(), zipfile.ZipFile(io.BytesIO(out)).read(sheet).decode()
    assert b.count("<x14:dataValidation ") == a.count("<x14:dataValidation ") == 3
    assert b.count("<dataValidation ") == a.count("<dataValidation ") == 763
    media = [n for n in zipfile.ZipFile(io.BytesIO(out)).namelist() if "media/" in n]
    assert len(media) == 22


def test_rows_written_and_reopen(produced):
    _, info, master, mapping, _, out = produced
    ws = openpyxl.load_workbook(io.BytesIO(out), data_only=True)["Template"]
    assert ws.max_row == info.data_row + len(master) - 1 == 615
    sku_col = next(c.index for c in info.columns if c.label == "SKU")
    assert ws.cell(info.data_row, sku_col).value == master.iloc[0]["Sku Code"]
    assert ws.cell(info.data_row + 607, sku_col).value == master.iloc[607]["Sku Code"]
    assert len(mapping) >= 8
