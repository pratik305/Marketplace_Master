"""Slice 6: can the tool's answer modes reproduce the hand-filled Amazon file, cell by cell?

For every column filled in `Amazon_Ring upd_ (1).xlsm` we look for the tool mode that explains it:
mapped from a master column, the same value for every row, or a value that depends on a master column.
The result goes through the real writer and is compared with the hand-filled file.
"""
import collections
import io
from pathlib import Path

import openpyxl
import pytest

from src.master_reader import read_master
from src.questions import Answer, resolve_all
from src.result import build_rows, clean
from src.template_reader import read_template
from src.validator import blocking, validate
from src.xlsm_writer import fill_template

SAMPLE = Path(__file__).resolve().parent.parent / "sample"
TEMPLATE, MASTER = SAMPLE / "RING (6).xlsm", SAMPLE / "Master_Arias_Ring_LOT3_v2.xlsx"
KEY = SAMPLE / "Amazon_Ring upd_ (1).xlsm"
pytestmark = pytest.mark.skipif(not (TEMPLATE.exists() and MASTER.exists() and KEY.exists()),
                                reason="sample/ files not present")


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


@pytest.fixture(scope="module")
def result():
    tb = TEMPLATE.read_bytes()
    info, master = read_template(tb), read_master(MASTER.read_bytes())
    ws = openpyxl.load_workbook(KEY, data_only=True)["Template"]
    n, r0 = len(master), info.data_row
    key = {c.index: [norm(ws.cell(r0 + i, c.index).value) for i in range(n)] for c in info.columns}
    mcols = {c: [norm(clean(v)) for v in master[c]] for c in master.columns}

    mapping, answers, unexplained = {}, {}, []
    for c in info.columns:
        vals = key[c.index]
        filled = [v for v in vals if v]
        if not filled:
            continue
        m = next((m for m, mv in mcols.items() if mv == vals), None)
        if m:
            mapping[c.index] = m
        elif len(set(filled)) == 1 and len(filled) == n:
            answers[c.index] = Answer("const", value=filled[0])
        else:
            for m, mv in mcols.items():
                if len(set(mv)) > 40:
                    continue
                table, ok = {}, True
                for a, b in zip(mv, vals):
                    if table.setdefault(a, b) != b:
                        ok = False
                        break
                if ok:
                    answers[c.index] = Answer("lookup", by=m, table={k: v for k, v in table.items() if v})
                    break
            else:
                unexplained.append(c.label)

    rows = build_rows(info, master, mapping, resolve_all(answers, master))
    out = fill_template(tb, r0, rows)
    wo = openpyxl.load_workbook(io.BytesIO(out), data_only=True)["Template"]
    total = match = 0
    diffs = collections.Counter()
    for c in info.columns:
        for i in range(n):
            k, o = key[c.index][i], norm(wo.cell(r0 + i, c.index).value)
            if k or o:
                total += 1
                match += k == o
                if k != o:
                    diffs[c.label] += 1
    return dict(info=info, rows=rows, mapping=mapping, answers=answers, unexplained=unexplained,
                total=total, match=match, diffs=diffs)


def test_most_columns_are_explained_by_the_three_modes(result):
    explained = len(result["mapping"]) + len(result["answers"])
    assert explained >= 70
    assert len(result["unexplained"]) <= 8      # image URLs and Material need a rule the PoC doesn't have
    print(f"\nexplained {explained} columns: {len(result['mapping'])} mapped, "
          f"{sum(a.mode == 'const' for a in result['answers'].values())} same-for-all, "
          f"{sum(a.mode == 'lookup' for a in result['answers'].values())} depends-on-a-column; "
          f"unexplained: {sorted(set(result['unexplained']))}")


def test_output_matches_the_hand_filled_file(result):
    share = result["match"] / result["total"]
    print(f"\ncells compared {result['total']}, matching {result['match']} ({share:.1%}); "
          f"differing columns: {result['diffs'].most_common(8)}")
    # The ~6% that differs is editorial, not a tool gap: the hand-filled file picked a different main image
    # (`_2.jpg` vs the master's `_1.jpg`), added photos that are not in the master, and filled Material.
    assert share >= 0.93
    assert set(result["diffs"]) <= set(result["unexplained"])      # only columns we could not explain differ


def test_output_has_no_blocking_problems(result):
    assert blocking(validate(result["info"], result["rows"])) == []
