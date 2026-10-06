"""Write values into an Amazon template without disturbing anything else in the file.

A normal Excel library rewrites the whole workbook and loses dropdown rules and images.
Here the .xlsm is opened as a zip, only the "Template" sheet's XML is replaced, and every
other entry is copied across byte-for-byte.
"""
import io
import re
import zipfile
from collections import Counter

from lxml import etree
from openpyxl.utils import get_column_letter

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")


class WriterError(Exception):
    """Raised with a plain message when the template can't be written."""


def _q(tag: str) -> str:
    return f"{{{NS}}}{tag}"


def sheet_path(zf: zipfile.ZipFile, sheet_name: str) -> str:
    """Find the zip path of a worksheet by its name, via workbook.xml and its relationships."""
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rid = None
    for sh in wb.iter(_q("sheet")):
        if sh.get("name") == sheet_name:
            rid = sh.get(f"{{{REL_NS}}}id")
    if rid is None:
        raise WriterError(f"Sheet '{sheet_name}' was not found in the template.")
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    for rel in rels.iter(f"{{{PKG_REL_NS}}}Relationship"):
        if rel.get("Id") == rid:
            target = rel.get("Target")
            return target.lstrip("/") if target.startswith("/") else "xl/" + target
    raise WriterError(f"Sheet '{sheet_name}' has no file in the template.")


def _cell(ref: str, value, style):
    c = etree.Element(_q("c"))
    c.set("r", ref)
    if style:
        c.set("s", style)
    if isinstance(value, bool):
        value = "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        v = etree.SubElement(c, _q("v"))
        v.text = str(int(value)) if float(value).is_integer() else repr(float(value))
    else:
        c.set("t", "inlineStr")
        is_ = etree.SubElement(c, _q("is"))
        t = etree.SubElement(is_, _q("t"))
        t.text = _ILLEGAL.sub("", str(value))
        t.set(f"{{{XML_NS}}}space", "preserve")
    return c


def fill_template(template_bytes: bytes, data_row: int, rows: list, sheet_name: str = "Template") -> bytes:
    """Return a copy of the template with `rows` written from `data_row` down.

    `rows` is a list of {column_index: value}. Empty values are left blank.
    """
    try:
        zin = zipfile.ZipFile(io.BytesIO(template_bytes))
    except zipfile.BadZipFile as exc:
        raise WriterError("The template is not a valid Excel file.") from exc
    path = sheet_path(zin, sheet_name)
    root = etree.fromstring(zin.read(path))
    sheet_data = root.find(_q("sheetData"))

    # Reuse the cell style of the first data row so the new cells look like the template's own.
    styles, common = {}, Counter()
    for row in sheet_data.findall(_q("row")):
        if int(row.get("r")) == data_row:
            for c in row.findall(_q("c")):
                col = re.match(r"[A-Z]+", c.get("r")).group(0)
                if c.get("s"):
                    styles[col] = c.get("s")
                    common[c.get("s")] += 1
    default_style = common.most_common(1)[0][0] if common else None

    for row in list(sheet_data.findall(_q("row"))):
        if int(row.get("r")) >= data_row:
            sheet_data.remove(row)

    max_col = 1
    for offset, values in enumerate(rows):
        r = data_row + offset
        row_el = etree.SubElement(sheet_data, _q("row"))
        row_el.set("r", str(r))
        for col_idx in sorted(values):
            value = values[col_idx]
            if value is None or value == "":
                continue
            letter = get_column_letter(col_idx)
            row_el.append(_cell(f"{letter}{r}", value, styles.get(letter, default_style)))
            max_col = max(max_col, col_idx)

    dim = root.find(_q("dimension"))
    if dim is not None and rows:
        m = re.match(r"([A-Z]+\d+):([A-Z]+)(\d+)", dim.get("ref") or "")
        if m:
            from openpyxl.utils import column_index_from_string
            last_col = max(column_index_from_string(m.group(2)), max_col)
            dim.set("ref", f"{m.group(1)}:{get_column_letter(last_col)}{max(int(m.group(3)), data_row + len(rows) - 1)}")

    new_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zout:
        for info in zin.infolist():
            data = new_xml if info.filename == path else zin.read(info.filename)
            zout.writestr(info, data, compress_type=info.compress_type)
    return out.getvalue()
