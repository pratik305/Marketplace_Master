"""Read an Amazon listing template (.xlsm): columns, required flags, dropdown values."""
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils import get_column_letter

REQUIRED = "Required"
CONDITIONAL = "Conditionally Required"


@dataclass
class TemplateColumn:
    index: int                 # 1-based column number on the Template sheet
    letter: str
    label: str                 # row 4, e.g. "Brand Name"
    key: str                   # row 5, e.g. "brand[marketplace_id=...]"
    group: str = ""            # row 3 group heading
    requirement: str = "Optional"
    description: str = ""
    example: str = ""
    allowed: list = field(default_factory=list)

    @property
    def required(self) -> bool:
        return self.requirement == REQUIRED

    @property
    def must_address(self) -> bool:
        return self.requirement in (REQUIRED, CONDITIONAL)

    @property
    def has_dropdown(self) -> bool:
        return bool(self.allowed)


@dataclass
class TemplateInfo:
    columns: list
    label_row: int
    key_row: int
    data_row: int
    defaults: dict = field(default_factory=dict)   # column index -> value pre-filled in the first data row

    def required_columns(self):
        return [c for c in self.columns if c.required]

    def must_address(self):
        """Required plus conditionally required: the columns a listing team has to deal with."""
        return [c for c in self.columns if c.must_address]

    def dropdown_columns(self):
        return [c for c in self.columns if c.has_dropdown]


class TemplateError(Exception):
    """Raised with a plain message when the template can't be understood."""


def _setting(settings: str, name: str, default: int) -> int:
    m = re.search(rf"(?:^|&){name}=(\d+)", settings or "")
    return int(m.group(1)) if m else default


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def read_template(data: bytes) -> TemplateInfo:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # corrupt or not an Excel file
        raise TemplateError(f"The template file could not be opened: {exc}") from exc
    for name in ("Template", "Data Definitions", "Valid Values"):
        if name not in wb.sheetnames:
            raise TemplateError(f"This does not look like an Amazon template: sheet '{name}' is missing.")

    ws = wb["Template"]
    settings = _text(ws.cell(1, 1).value)
    label_row = _setting(settings, "labelRow", 4)
    key_row = _setting(settings, "attributeRow", 5)
    data_row = _setting(settings, "dataRow", 8)
    group_row = label_row - 1

    # Required flags and descriptions come from "Data Definitions", joined by attribute key.
    # Each row is: attribute key, label, accepted values text, example, requirement.
    # The sheets start in column B, so use the first filled cell of the row rather than a fixed column.
    defs = {}
    for row in wb["Data Definitions"].iter_rows(min_row=3, values_only=True):
        cells = [_text(v) for v in row]
        start = next((i for i, v in enumerate(cells) if v), None)
        if start is not None and len(cells) >= start + 5 and cells[start + 4]:
            defs[cells[start]] = (cells[start + 4], cells[start + 2], cells[start + 3])

    # Allowed dropdown values come from "Valid Values", rows labelled "Field - [ PRODUCT_TYPE ]".
    valid = defaultdict(list)       # label -> list of value lists, in sheet order
    for row in wb["Valid Values"].iter_rows(values_only=True):
        cells = [_text(v) for v in row]
        start = next((i for i, v in enumerate(cells) if v), None)
        if start is None:
            continue
        m = re.match(r"^(.*?)\s+-\s+\[.*\]\s*$", cells[start])
        if m:
            valid[m.group(1).strip()].append([v for v in cells[start + 1:] if v])

    columns, seen, group = [], defaultdict(int), ""
    for idx in range(1, ws.max_column + 1):
        label = _text(ws.cell(label_row, idx).value)
        key = _text(ws.cell(key_row, idx).value)
        g = _text(ws.cell(group_row, idx).value)
        group = g or group
        if not label and not key:
            continue
        req, desc, example = defs.get(key, ("Optional", "", ""))
        n = seen[label]
        seen[label] += 1
        lists = valid.get(label, [])
        allowed = lists[n] if n < len(lists) else []
        columns.append(TemplateColumn(
            index=idx, letter=get_column_letter(idx), label=label, key=key,
            group=group, requirement=req, description=desc, example=example, allowed=allowed))
    if not columns:
        raise TemplateError("No columns were found in the Template sheet.")
    defaults = {c.index: ws.cell(data_row, c.index).value for c in columns
                if ws.cell(data_row, c.index).value not in (None, "")}
    return TemplateInfo(columns=columns, label_row=label_row, key_row=key_row,
                        data_row=data_row, defaults=defaults)
